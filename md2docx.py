#!/usr/bin/env python3
"""Convierte archivos Markdown a documentos Word (.docx).

Uso:
    python md2docx.py entrada.md [-o salida.docx]
"""

import argparse
import io
import re
import sys
import urllib.request
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Emu, Inches, Pt, RGBColor

HARD_BREAK = "\x00"
MAX_IMAGE_WIDTH = Inches(6)

TOKEN_RE = re.compile(
    r"(?P<ticks>`+)(?P<code_body>.+?)(?P=ticks)"
    r"|!\[(?P<image_alt>[^\]]*)\]\((?P<image_src>(?:[^()\s]|\([^()\s]*\))+)(?:\s+\"[^\"]*\")?\)"
    r"|\[(?P<link_text>[^\]]+)\]\((?P<link_href>(?:[^()\s]|\([^()\s]*\))+)(?:\s+\"[^\"]*\")?\)"
    r"|(?P<strongem1>\*\*\*(?P<se1>.+?)\*\*\*)"
    r"|(?P<strongem2>___(?P<se2>.+?)___)"
    r"|(?P<strong1>\*\*(?P<s1>.+?)\*\*)"
    r"|(?P<strong2>__(?P<s2>.+?)__)"
    r"|(?P<em1>(?<!\*)\*(?P<e1>[^*\n]+)\*(?!\*))"
    r"|(?P<em2>(?<!\w)_(?P<e2>[^_\n]+)_(?!\w))"
    r"|(?P<strike>~~(?P<st>.+?)~~)"
)

FENCE_RE = re.compile(r"^\s*(`{3,}|~{3,})\s*([\w+#-]*)\s*$")
HEADING_RE = re.compile(r"^(#{1,6})\s+(.*?)\s*#*\s*$")
HR_RE = re.compile(r"^\s{0,3}([-*_])(?:\s*\1){2,}\s*$")
LIST_RE = re.compile(r"^(\s*)([-*+]|\d+[.)])\s+(.*)$")
TASK_RE = re.compile(r"^\[([ xX])\]\s+(.*)$")
SETEXT_H1_RE = re.compile(r"^=+\s*$")
SETEXT_H2_RE = re.compile(r"^-{2,}\s*$")
TABLE_SEP_RE = re.compile(r"^\s*\|?\s*:?-{1,}:?\s*(\|\s*:?-{1,}:?\s*)*\|?\s*$")
ESCAPE_RE = re.compile(r"\\([\\`*_{}\[\]()#+\-.!|~>])")

BULLETS = ["•", "◦", "▪", "•", "◦", "▪", "•", "◦", "▪"]

BASE_DIR = Path.cwd()


def protect_escapes(text):
    table = {}

    def repl(m):
        key = "\x01%d\x01" % len(table)
        table[key] = m.group(1)
        return key

    return ESCAPE_RE.sub(repl, text), table


def restore(text, table):
    for key, value in table.items():
        text = text.replace(key, value)
    return text


def set_rfonts(run, name):
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rFonts.set(qn(attr), name)


def add_hyperlink(paragraph, text, url):
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rPr.append(color)
    rPr.append(underline)
    run.append(rPr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def fetch_image(src):
    if re.match(r"^https?://", src):
        req = urllib.request.Request(src, headers={"User-Agent": "md2docx/1.0"})
        with urllib.request.urlopen(req, timeout=15) as resp:
            return resp.read()
    candidates = [Path(src).expanduser()]
    if not Path(src).is_absolute():
        candidates.insert(0, BASE_DIR / src)
    for path in candidates:
        if path.is_file():
            return path.read_bytes()
    return None


def add_image(paragraph, src, alt):
    try:
        data = fetch_image(src)
    except Exception:
        data = None
    if not data:
        if alt:
            run = paragraph.add_run(alt)
            run.italic = True
        return
    stream = io.BytesIO(data)
    try:
        from docx.image.image import Image as DocxImage

        native = DocxImage.from_file(io.BytesIO(data)).width
        width = native if native and native < MAX_IMAGE_WIDTH else MAX_IMAGE_WIDTH
        paragraph.add_run().add_picture(stream, width=Emu(int(width)))
    except Exception:
        if alt:
            run = paragraph.add_run(alt)
            run.italic = True


def plain_run(paragraph, text, bold, italic, esc):
    if not text:
        return
    run = paragraph.add_run(restore(text, esc))
    if bold:
        run.bold = True
    if italic:
        run.italic = True


def tokenize(paragraph, text, bold=False, italic=False, esc=None):
    if esc is None:
        text, esc = protect_escapes(text)
    pos = 0
    for m in TOKEN_RE.finditer(text):
        if m.start() > pos:
            plain_run(paragraph, text[pos : m.start()], bold, italic, esc)
        g = m.groupdict()
        if g["ticks"] is not None:
            run = paragraph.add_run(restore(g["code_body"], esc))
            run.font.name = "Consolas"
            run.font.size = Pt(10)
            set_rfonts(run, "Consolas")
        elif g["image_alt"] is not None:
            add_image(paragraph, g["image_src"], restore(g["image_alt"], esc))
        elif g["link_text"] is not None:
            add_hyperlink(paragraph, restore(g["link_text"], esc), g["link_href"])
        elif g["se1"] is not None or g["se2"] is not None:
            add_runs(paragraph, g["se1"] or g["se2"], bold=True, italic=True, esc=esc)
        elif g["s1"] is not None or g["s2"] is not None:
            add_runs(paragraph, g["s1"] or g["s2"], bold=True, italic=italic, esc=esc)
        elif g["e1"] is not None or g["e2"] is not None:
            add_runs(paragraph, g["e1"] or g["e2"], bold=bold, italic=True, esc=esc)
        elif g["st"] is not None:
            before = len(paragraph.runs)
            add_runs(paragraph, g["st"], bold=bold, italic=italic, esc=esc)
            for run in paragraph.runs[before:]:
                run.font.strike = True
        pos = m.end()
    if pos < len(text):
        plain_run(paragraph, text[pos:], bold, italic, esc)


def add_runs(paragraph, text, bold=False, italic=False, esc=None):
    if esc is None:
        text, esc = protect_escapes(text)
    parts = text.split(HARD_BREAK)
    for i, part in enumerate(parts):
        if i:
            paragraph.add_run().add_break()
        tokenize(paragraph, part, bold=bold, italic=italic, esc=esc)


def ensure_code_style(doc):
    styles = doc.styles
    try:
        style = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    except Exception:
        style = styles["Code Block"]
    style.base_style = styles["Normal"]
    style.font.name = "Consolas"
    style.font.size = Pt(9.5)
    style.font.color.rgb = RGBColor(0x24, 0x29, 0x2E)
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rFonts.set(qn(attr), "Consolas")
    return style


def shade_paragraph(paragraph, fill):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def add_code_block(doc, code, lang):
    paragraph = doc.add_paragraph(style="Code Block")
    if lang:
        paragraph.paragraph_format.space_before = Pt(6)
    lines = code.split("\n")
    while lines and not lines[-1].strip():
        lines.pop()
    for i, line in enumerate(lines):
        if i:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line if line else " ")
        set_rfonts(run, "Consolas")
    shade_paragraph(paragraph, "F6F8FA")
    return paragraph


def add_horizontal_rule(doc):
    paragraph = doc.add_paragraph()
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "A6A6A6")
    pBdr.append(bottom)
    pPr.append(pBdr)


class Numbering:
    def __init__(self, doc):
        self.root = doc.part.numbering_part.element
        abstract_ids = [
            int(e.get(qn("w:abstractNumId")))
            for e in self.root.findall(qn("w:abstractNum"))
        ]
        num_ids = [int(e.get(qn("w:numId"))) for e in self.root.findall(qn("w:num"))]
        self._next_abstract = max(abstract_ids, default=0) + 1
        self._next_num = max(num_ids, default=0) + 1
        self.bullet_abstract = self._add_abstract("bullet")
        self.decimal_abstract = self._add_abstract("decimal")
        self.bullet_num = self.new_num(self.bullet_abstract)

    def _insert_index(self):
        first_num = self.root.find(qn("w:num"))
        if first_num is None:
            return len(self.root)
        return list(self.root).index(first_num)

    def _add_abstract(self, num_fmt):
        abstract_id = self._next_abstract
        self._next_abstract += 1
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "multilevel")
        abstract.append(multi)
        for level in range(9):
            lvl = OxmlElement("w:lvl")
            lvl.set(qn("w:ilvl"), str(level))
            start = OxmlElement("w:start")
            start.set(qn("w:val"), "1")
            fmt = OxmlElement("w:numFmt")
            fmt.set(qn("w:val"), num_fmt)
            text = OxmlElement("w:lvlText")
            if num_fmt == "bullet":
                text.set(qn("w:val"), BULLETS[level])
            else:
                text.set(qn("w:val"), "%%%d." % (level + 1))
            jc = OxmlElement("w:lvlJc")
            jc.set(qn("w:val"), "left")
            pPr = OxmlElement("w:pPr")
            ind = OxmlElement("w:ind")
            ind.set(qn("w:left"), str(720 * (level + 1)))
            ind.set(qn("w:hanging"), "360")
            pPr.append(ind)
            lvl.append(start)
            lvl.append(fmt)
            lvl.append(text)
            lvl.append(jc)
            lvl.append(pPr)
            abstract.append(lvl)
        self.root.insert(self._insert_index(), abstract)
        return abstract_id

    def new_num(self, abstract_id):
        num_id = self._next_num
        self._next_num += 1
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        ref = OxmlElement("w:abstractNumId")
        ref.set(qn("w:val"), str(abstract_id))
        num.append(ref)
        self.root.append(num)
        return num_id

    @staticmethod
    def apply(paragraph, num_id, level):
        pPr = paragraph._p.get_or_add_pPr()
        numPr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), str(level))
        num = OxmlElement("w:numId")
        num.set(qn("w:val"), str(num_id))
        numPr.append(ilvl)
        numPr.append(num)
        pStyle = pPr.find(qn("w:pStyle"))
        if pStyle is not None:
            pPr.insert(list(pPr).index(pStyle) + 1, numPr)
        else:
            pPr.insert(0, numPr)


def split_table_row(line):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    cells = re.split(r"(?<!\\)\|", line)
    return [c.replace("\\|", "|").strip() for c in cells]


def parse_aligns(line):
    aligns = []
    for cell in split_table_row(line):
        cell = cell.strip()
        if cell.startswith(":") and cell.endswith(":"):
            aligns.append(WD_ALIGN_PARAGRAPH.CENTER)
        elif cell.endswith(":"):
            aligns.append(WD_ALIGN_PARAGRAPH.RIGHT)
        elif cell.startswith(":"):
            aligns.append(WD_ALIGN_PARAGRAPH.LEFT)
        else:
            aligns.append(None)
    return aligns


def is_list_item(line):
    return bool(LIST_RE.match(line))


def parse_list_block(lines, i):
    items = []
    n = len(lines)
    base_ordered = None
    while i < n:
        line = lines[i]
        if is_list_item(line):
            m = LIST_RE.match(line)
            indent = len(m.group(1).expandtabs(4))
            ordered = m.group(2)[0].isdigit()
            if indent == 0:
                if base_ordered is None:
                    base_ordered = ordered
                elif ordered != base_ordered:
                    break
            items.append(line)
            i += 1
        elif not line.strip():
            j = i
            while j < n and not lines[j].strip():
                j += 1
            if j < n and is_list_item(lines[j]):
                i = j
            else:
                break
        else:
            break
    return items, i


def render_list(doc, numbering, items):
    block_num = numbering.new_num(numbering.decimal_abstract)
    stack = [0]
    for raw in items:
        m = LIST_RE.match(raw)
        indent = len(m.group(1).expandtabs(4))
        marker = m.group(2)
        content = m.group(3)
        if indent > stack[-1]:
            stack.append(indent)
        elif indent < stack[-1]:
            while len(stack) > 1 and indent < stack[-1]:
                stack.pop()
            if indent != stack[-1]:
                stack.append(indent)
        level = min(len(stack) - 1, 8)
        ordered = marker[0].isdigit()
        task = TASK_RE.match(content)
        paragraph = doc.add_paragraph()
        if task:
            checked = task.group(1).lower() == "x"
            box = paragraph.add_run("☑ " if checked else "☐ ")
            box.font.name = "Segoe UI Symbol"
            set_rfonts(box, "Segoe UI Symbol")
            content = task.group(2)
        add_runs(paragraph, content)
        Numbering.apply(
            paragraph,
            block_num if ordered else numbering.bullet_num,
            level,
        )


def render_table(doc, header, sep, body):
    columns = len(header)
    aligns = parse_aligns(sep)
    table = doc.add_table(rows=1 + len(body), cols=columns)
    table.style = "Table Grid"
    for col, text in enumerate(header):
        cell = table.rows[0].cells[col]
        cell.paragraphs[0].clear()
        add_runs(cell.paragraphs[0], text, bold=True)
        if col < len(aligns) and aligns[col]:
            cell.paragraphs[0].alignment = aligns[col]
    for row_idx, row in enumerate(body, start=1):
        cells = row if len(row) == columns else (row + [""] * columns)[:columns]
        for col, text in enumerate(cells):
            cell = table.rows[row_idx].cells[col]
            cell.paragraphs[0].clear()
            add_runs(cell.paragraphs[0], text)
            if col < len(aligns) and aligns[col]:
                cell.paragraphs[0].alignment = aligns[col]


def strip_front_matter(lines):
    if lines and lines[0].strip() == "---":
        for i in range(1, len(lines)):
            if lines[i].strip() in ("---", "..."):
                return lines[i + 1 :]
    return lines


def convert(markdown, doc):
    ensure_code_style(doc)
    numbering = Numbering(doc)
    lines = strip_front_matter(markdown.replace("\r\n", "\n").replace("\r", "\n").split("\n"))
    n = len(lines)
    i = 0
    while i < n:
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        fence = FENCE_RE.match(line)
        if fence:
            marker = fence.group(1)
            lang = fence.group(2)
            i += 1
            body = []
            while i < n and not lines[i].strip().startswith(marker[:3]):
                body.append(lines[i])
                i += 1
            i += 1
            add_code_block(doc, "\n".join(body), lang)
            continue

        if HR_RE.match(line):
            add_horizontal_rule(doc)
            i += 1
            continue

        heading = HEADING_RE.match(line)
        if heading:
            level = len(heading.group(1))
            paragraph = doc.add_heading(level=min(level, 9))
            add_runs(paragraph, heading.group(2))
            i += 1
            continue

        if is_list_item(line):
            items, i = parse_list_block(lines, i)
            render_list(doc, numbering, items)
            continue

        if line.lstrip().startswith(">"):
            quote_lines = []
            while i < n and lines[i].lstrip().startswith(">"):
                quote_lines.append(re.sub(r"^\s*>\s?", "", lines[i]))
                i += 1
            for block in re.split(r"\n\s*\n", "\n".join(quote_lines)):
                if not block.strip():
                    continue
                paragraph = doc.add_paragraph(style="Quote")
                add_runs(paragraph, " ".join(block.split()))
            continue

        if "|" in line and i + 1 < n and TABLE_SEP_RE.match(lines[i + 1]) and "|" in lines[i + 1]:
            header = split_table_row(line)
            sep = lines[i + 1]
            i += 2
            body = []
            while i < n and lines[i].strip() and "|" in lines[i]:
                body.append(split_table_row(lines[i]))
                i += 1
            render_table(doc, header, sep, body)
            continue

        paragraph_lines = [line]
        i += 1
        while i < n and lines[i].strip():
            nxt = lines[i]
            if SETEXT_H1_RE.match(nxt) or SETEXT_H2_RE.match(nxt):
                paragraph_lines.append(nxt)
                i += 1
                break
            if (
                HEADING_RE.match(nxt)
                or HR_RE.match(nxt)
                or FENCE_RE.match(nxt)
                or is_list_item(nxt)
                or nxt.lstrip().startswith(">")
                or ("|" in nxt and i + 1 < n and TABLE_SEP_RE.match(lines[i + 1]))
            ):
                break
            paragraph_lines.append(nxt)
            i += 1

        last = paragraph_lines[-1]
        if len(paragraph_lines) > 1 and SETEXT_H1_RE.match(last):
            paragraph = doc.add_heading(level=1)
            paragraph_lines.pop()
        elif len(paragraph_lines) > 1 and SETEXT_H2_RE.match(last):
            paragraph = doc.add_heading(level=2)
            paragraph_lines.pop()
        else:
            paragraph = doc.add_paragraph()

        parts = []
        for idx, pline in enumerate(paragraph_lines):
            trailing = len(pline) - len(pline.rstrip())
            text = pline.strip()
            parts.append(text)
            if idx < len(paragraph_lines) - 1:
                parts.append(HARD_BREAK if trailing >= 2 else " ")
        add_runs(paragraph, "".join(parts))


def main():
    parser = argparse.ArgumentParser(description="Convierte Markdown a Word (.docx)")
    parser.add_argument("input", help="archivo Markdown de entrada")
    parser.add_argument("-o", "--output", help="archivo .docx de salida")
    parser.add_argument("--title", help="título del documento")
    args = parser.parse_args()

    input_path = Path(args.input)
    if not input_path.is_file():
        parser.error("no existe el archivo: %s" % input_path)

    output_path = Path(args.output) if args.output else input_path.with_suffix(".docx")

    markdown = input_path.read_text(encoding="utf-8")

    global BASE_DIR
    BASE_DIR = input_path.resolve().parent

    doc = Document()
    first_heading = re.search(r"^#{1}\s+(.+)$", markdown, re.M)
    doc.core_properties.title = (
        args.title or (first_heading.group(1).strip() if first_heading else input_path.stem)
    )

    convert(markdown, doc)
    doc.save(str(output_path))
    print("Generado: %s" % output_path)


if __name__ == "__main__":
    sys.exit(main())
